"""Apply humanized rewrites back into the original .docx templates.

Each rewrite module exports a list `BLOCKS` of (style_name, text) tuples.
Style names map to the Word styles already defined in the original document
("Heading 1", "Heading 2", "Heading 3", "Normal", "Body Text"). Empty text
with style "" inserts a blank paragraph.

We open the original doc, remove all body paragraphs/tables, then append the
new content using the existing styles. Sections, page setup, headers/footers,
and style definitions inherited from the original are preserved.
"""
from __future__ import annotations

import importlib
import sys
from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn


def clear_body(doc):
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is sectPr:
            continue
        body.remove(child)


def add_block(doc, style: str, text: str):
    if style == "":
        p = doc.add_paragraph("")
        return p
    p = doc.add_paragraph(text, style=style)
    return p


def rebuild(src: Path, dst: Path, blocks):
    doc = Document(str(src))
    clear_body(doc)
    for style, text in blocks:
        add_block(doc, style, text)
    doc.save(str(dst))


JOBS = [
    (
        "humanize_rmk4",
        "RMK/01079_Dzaki Muhammad Yusfian_RMK Pert. 4.docx",
    ),
    (
        "humanize_rmk5",
        "RMK/01079_Dzaki Muhammad Yusfian_RMK Pert. 5.docx",
    ),
    (
        "humanize_art7",
        "Critical Thinking of the Article/01079_Dzaki Muhammad Yusfian_Artikel 7.docx",
    ),
    (
        "humanize_art8",
        "Critical Thinking of the Article/01079_Dzaki Muhammad Yusfian_Artikel 8.docx",
    ),
]


def main(only: str | None = None):
    repo = Path(__file__).resolve().parents[2]
    sys.path.insert(0, str(Path(__file__).parent))
    for module_name, rel_path in JOBS:
        if only and only not in module_name:
            continue
        target = repo / rel_path
        mod = importlib.import_module(module_name)
        blocks = mod.BLOCKS
        rebuild(target, target, blocks)
        print(f"OK  {rel_path}  ({len(blocks)} blocks)")


if __name__ == "__main__":
    arg = sys.argv[1] if len(sys.argv) > 1 else None
    main(arg)
