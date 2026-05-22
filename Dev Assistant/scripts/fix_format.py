"""Apply two formatting fixes to RMK Pra UTS.docx:

1. Justify every body paragraph (Normal, Body Text, First Paragraph styles).
   Headings stay left-aligned; Captions stay centered (under images).
2. Force visible grid borders on every table by setting Table Grid style
   AND writing explicit XML border definitions so borders render in Word
   regardless of theme.
"""
from __future__ import annotations
import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "RMK Pra UTS.docx"

BODY_STYLES = {"Normal", "Body Text", "First Paragraph"}
SKIP_STYLES = {"Caption"}  # leave centered

BORDER_KINDS = ("top", "left", "bottom", "right", "insideH", "insideV")


def set_table_borders(table) -> None:
    """Add a <w:tblBorders> with single black borders to the table XML."""
    tbl = table._element
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    # Remove any existing borders to start fresh
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for kind in BORDER_KINDS:
        b = OxmlElement(f"w:{kind}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "6")        # 0.75 pt
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tblPr.append(borders)


def set_cell_borders(cell) -> None:
    """Ensure each cell has explicit single black borders on all 4 sides."""
    tcPr = cell._tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        cell._tc.insert(0, tcPr)
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    tcBorders = OxmlElement("w:tcBorders")
    for kind in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{kind}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "6")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        tcBorders.append(b)
    tcPr.append(tcBorders)


def main() -> int:
    if not DOCX.exists():
        print(f"docx not found: {DOCX}", file=sys.stderr)
        return 1

    doc = Document(DOCX)

    # 1. Justify body paragraphs
    n_justified = 0
    for p in doc.paragraphs:
        style_name = p.style.name
        if style_name in SKIP_STYLES:
            continue
        if style_name.startswith("Heading"):
            continue
        if style_name in BODY_STYLES:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            n_justified += 1

    # Also justify paragraphs inside table cells (these are usually Normal style)
    n_cell_justified = 0
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.style.name not in SKIP_STYLES and not p.style.name.startswith("Heading"):
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        n_cell_justified += 1

    # 2. Apply Table Grid style + explicit borders to every table and cell
    n_tables = 0
    for tbl in doc.tables:
        try:
            tbl.style = "Table Grid"
        except KeyError:
            pass  # style not defined in this doc; XML borders below still work
        set_table_borders(tbl)
        for row in tbl.rows:
            for cell in row.cells:
                set_cell_borders(cell)
        n_tables += 1

    doc.save(DOCX)
    print(f"justified {n_justified} body paragraphs, "
          f"{n_cell_justified} table-cell paragraphs")
    print(f"applied grid borders to {n_tables} tables")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
