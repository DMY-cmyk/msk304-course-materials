"""Verify Pertemuan 1 addition meets all acceptance criteria from the spec.

Exits 0 on success, nonzero on first failure with a descriptive message.
"""
from __future__ import annotations
import sys
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "RMK Pra UTS.docx"

EXPECTED_H2_HEADINGS = [
    "§1.1 Pengantar & Pemetaan Bab",
    "§1.2 Apa Itu Strategi? — Lima Tindakan Inti",
    "§1.3 Strategi vs Model Bisnis",
    "§1.4 Keunggulan Kompetitif yang Berkelanjutan",
    "§1.5 Kapabilitas Dinamis & Evolusi Strategi",
    "§1.6 Tiga Uji Strategi yang Menang",
    "§1.7 Indonesian Flagship — BCA Lulus Tiga Uji",
    "§1.8 Sintesis Pertemuan 1 dan Jembatan ke Pertemuan 2",
]


def fail(msg: str) -> int:
    print(f"FAIL: {msg}", file=sys.stderr)
    return 1


def main() -> int:
    if not DOCX.exists():
        return fail(f"docx not found: {DOCX}")
    doc = Document(DOCX)

    # 1. Pertemuan 1 Heading 1 exists exactly once
    h1_pert1 = [p for p in doc.paragraphs
                if p.style.name == "Heading 1" and "PERTEMUAN 1" in p.text.upper()]
    if not h1_pert1:
        return fail("no Heading 1 containing 'PERTEMUAN 1' found")
    if len(h1_pert1) > 1:
        return fail(f"expected exactly 1 Pertemuan 1 H1, got {len(h1_pert1)}")

    # 2. All expected §1.x H2 sub-headings present
    h2_texts = [p.text.strip() for p in doc.paragraphs
                if p.style.name == "Heading 2"]
    for expected in EXPECTED_H2_HEADINGS:
        if expected not in h2_texts:
            return fail(f"missing H2: {expected!r}")

    # 3. Table count is exactly 4 (was 2, plus 2 new)
    if len(doc.tables) != 4:
        return fail(f"expected 4 tables, got {len(doc.tables)}")

    # 4. BAGIAN I §1.1 title updated
    found_updated = any(
        "§1.1 Alur Naratif Induk (Pertemuan 1 → 7)" in p.text
        for p in doc.paragraphs
    )
    if not found_updated:
        return fail("BAGIAN I §1.1 title not updated to '(Pertemuan 1 → 7)'")

    # 5. At least 3 inline images
    n_images = len(doc.inline_shapes)
    if n_images < 3:
        return fail(f"expected at least 3 inline images, got {n_images}")

    # 6. Paragraph count in expected band
    n_paras = len(doc.paragraphs)
    if not (990 <= n_paras <= 1010):
        print(f"WARN: paragraph count {n_paras} outside expected [990, 1010]",
              file=sys.stderr)

    print(f"PASS — {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, "
          f"{n_images} inline images")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
