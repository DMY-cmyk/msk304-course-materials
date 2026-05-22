"""Final acceptance verification for Pertemuan 2-7 visuals addition.

Exits 0 on PASS. Exits 1 on any failure with descriptive message.
"""
from __future__ import annotations
import sys
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "RMK Pra UTS.docx"

EXPECTED_GAMBAR_PER_PERT = {2: 6, 3: 8, 4: 3, 5: 7, 6: 4, 7: 5}
EXPECTED_TABEL_PER_PERT  = {2: 4, 3: 3, 4: 2, 5: 0, 6: 0, 7: 2}


def fail(msg: str) -> int:
    print(f"FAIL: {msg}", file=sys.stderr)
    return 1


def main() -> int:
    if not DOCX.exists():
        return fail(f"docx not found: {DOCX}")
    doc = Document(DOCX)

    n_imgs = len(doc.inline_shapes)
    if n_imgs < 36:
        return fail(f"expected >=36 inline images, got {n_imgs}")

    n_tables = len(doc.tables)
    if n_tables != 15:
        return fail(f"expected 15 tables, got {n_tables}")

    gambar_counts = {n: 0 for n in EXPECTED_GAMBAR_PER_PERT}
    tabel_counts  = {n: 0 for n in EXPECTED_TABEL_PER_PERT}
    for p in doc.paragraphs:
        if p.style.name != "Caption":
            continue
        t = p.text.strip()
        for n in gambar_counts:
            if t.startswith(f"Gambar {n}."):
                gambar_counts[n] += 1
        for n in tabel_counts:
            if t.startswith(f"Tabel {n}."):
                tabel_counts[n] += 1

    for n, expected in EXPECTED_GAMBAR_PER_PERT.items():
        got = gambar_counts[n]
        if got != expected:
            return fail(f"Pertemuan {n}: expected {expected} 'Gambar {n}.x' "
                        f"captions, got {got}")

    for n, expected in EXPECTED_TABEL_PER_PERT.items():
        got = tabel_counts[n]
        if got != expected:
            return fail(f"Pertemuan {n}: expected {expected} 'Tabel {n}.x' "
                        f"captions, got {got}")

    pert1_imgs = sum(1 for p in doc.paragraphs
                     if p.style.name == "Caption" and p.text.startswith("Gambar 1."))
    if pert1_imgs != 3:
        return fail(f"Pertemuan 1 regression: expected 3 'Gambar 1.x' captions, "
                    f"got {pert1_imgs}")

    print(f"PASS — {n_imgs} inline images, {n_tables} tables.")
    print("Per-Pertemuan caption counts:")
    for n in sorted(EXPECTED_GAMBAR_PER_PERT):
        print(f"  Pertemuan {n}: {gambar_counts[n]} Gambar, "
              f"{tabel_counts[n]} Tabel")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
