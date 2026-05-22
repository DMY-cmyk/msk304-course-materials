"""Post-build sanity checks on the final DOCX."""
from __future__ import annotations
import sys
from pathlib import Path

from docx import Document


def verify(docx_path: Path) -> list[str]:
    failures: list[str] = []
    if not docx_path.exists():
        return [f"DOCX not found: {docx_path}"]

    doc = Document(str(docx_path))

    # Check 5: no surviving sentinels
    all_text = "\n".join(p.text for p in doc.paragraphs)
    if "§T" in all_text:
        failures.append("Found surviving §T sentinels in output")

    # Check 4: at least 10 images embedded (via relationship count)
    rels = doc.part.rels
    img_rel_count = sum(1 for r in rels.values() if "image" in r.target_ref)
    if img_rel_count < 10:
        failures.append(f"Only {img_rel_count} images embedded (expected >=10 for 7 chapters)")

    # Check 6: at least 7 Heading 1 paragraphs (one per chapter)
    h1_count = sum(1 for p in doc.paragraphs if p.style and p.style.name == "Heading 1")
    if h1_count < 7:
        failures.append(f"Only {h1_count} Heading 1 paragraphs (expected >=7, one per chapter)")

    return failures


if __name__ == "__main__":
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Temu 1-7 Buku Gamble (Bahasa Indonesia).docx")
    fails = verify(path)
    if fails:
        for f in fails:
            print(f"FAIL: {f}")
        sys.exit(1)
    print("OK")
