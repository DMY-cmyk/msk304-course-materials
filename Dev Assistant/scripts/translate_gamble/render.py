"""Stage 3: translated JSON -> single DOCX with italic terms + embedded images."""
from __future__ import annotations
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt


def italic_runs(text: str, spans: list[tuple[int, int]]) -> list[tuple[str, bool]]:
    """Split text into (substring, is_italic) runs at the given char spans."""
    if not spans:
        return [(text, False)]
    spans_sorted = sorted(spans)
    runs: list[tuple[str, bool]] = []
    cursor = 0
    for start, end in spans_sorted:
        if start < cursor:
            continue  # skip overlaps
        if start > cursor:
            runs.append((text[cursor:start], False))
        runs.append((text[start:end], True))
        cursor = end
    if cursor < len(text):
        runs.append((text[cursor:], False))
    return runs


def _add_paragraph_with_runs(
    doc,
    text: str,
    spans: list[tuple[int, int]],
    *,
    style: str | None = None,
    align: int | None = None,
    italic_all: bool = False,
) -> None:
    para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        para.alignment = align
    for substring, is_italic in italic_runs(text, spans):
        run = para.add_run(substring)
        run.italic = is_italic or italic_all
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)


def _add_heading(doc, text: str, level: int, page_break: bool) -> None:
    if page_break:
        br = doc.add_paragraph()
        br.add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_heading(text, level=max(1, min(level, 3)))
    for run in h.runs:
        run.font.name = "Times New Roman"


def render_to_docx(translated_pages_dir: Path, images_dir: Path, out_path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    page_files = sorted(translated_pages_dir.glob("page_*.json"))
    last_chapter: int | None = None

    for src in page_files:
        payload: dict[str, Any] = json.loads(src.read_text(encoding="utf-8"))
        chapter = payload.get("chapter")
        for block in payload["blocks"]:
            btype = block["type"]
            spans = [tuple(p) for p in block.get("italic_spans", [])]
            text = block.get("text", "")

            if btype == "heading":
                level = int(block.get("level", 2))
                new_chapter = (
                    level == 1
                    and chapter is not None
                    and chapter != last_chapter
                )
                _add_heading(doc, text, level, page_break=(new_chapter and last_chapter is not None))
                if new_chapter:
                    last_chapter = chapter
            elif btype == "body":
                _add_paragraph_with_runs(doc, text, spans, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            elif btype == "list_item":
                _add_paragraph_with_runs(doc, text, spans, style="List Bullet", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            elif btype == "caption":
                _add_paragraph_with_runs(doc, text, spans, align=WD_ALIGN_PARAGRAPH.CENTER, italic_all=True)
            elif btype == "image":
                ref = block.get("ref")
                if not ref:
                    continue
                img_path = images_dir / ref
                if not img_path.exists():
                    continue
                doc.add_picture(str(img_path), width=Inches(5.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(out_path))


if __name__ == "__main__":
    import sys
    base = Path("Dev Assistant/scripts/translate_gamble/build")
    pages = Path(sys.argv[1]) if len(sys.argv) > 1 else base / "translated"
    imgs = Path(sys.argv[2]) if len(sys.argv) > 2 else base / "extracted" / "images"
    out = Path(sys.argv[3]) if len(sys.argv) > 3 else Path("Temu 1-7 Buku Gamble (Bahasa Indonesia).docx")
    render_to_docx(pages, imgs, out)
    print(f"Wrote {out}")
