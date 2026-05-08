"""Extract paragraph text and basic structure from .docx files for analysis."""
import sys
from pathlib import Path
from docx import Document


def extract(path: Path) -> str:
    doc = Document(str(path))
    out = [f"===== FILE: {path.name} =====\n"]
    for i, p in enumerate(doc.paragraphs):
        style = p.style.name if p.style else ""
        text = p.text
        if not text.strip() and style.lower() == "normal":
            out.append("")
            continue
        out.append(f"[{style}] {text}")
    # tables
    for ti, table in enumerate(doc.tables):
        out.append(f"\n--- TABLE {ti} ---")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
            out.append(" || ".join(cells))
    return "\n".join(out)


if __name__ == "__main__":
    for p in sys.argv[1:]:
        print(extract(Path(p)))
        print("\n\n")
