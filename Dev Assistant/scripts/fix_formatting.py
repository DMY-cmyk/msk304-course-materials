# -*- coding: utf-8 -*-
"""Fix page margins (A4, Left 4 / Right 3 / Top 3 / Bottom 3 cm)
and paragraph alignment (Justify for body; Left for headings) on all
submission docx files.

Run from the repo root:
    python "Dev Assistant/scripts/fix_formatting.py"
"""
from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ── Target page setup ────────────────────────────────────────────────────────
# python-docx stores lengths in EMU (English Metric Units)
# 1 inch = 914400 EMU; 1 cm = 914400 / 2.54
_CM = 914400 / 2.54

A4_W   = int(21.0  * _CM)
A4_H   = int(29.7  * _CM)
M_TOP  = int(3.0   * _CM)
M_BOT  = int(3.0   * _CM)
M_LEFT = int(4.0   * _CM)
M_RGT  = int(3.0   * _CM)

# ── Style → alignment mapping ────────────────────────────────────────────────
JUSTIFY   = WD_ALIGN_PARAGRAPH.JUSTIFY
LEFT      = WD_ALIGN_PARAGRAPH.LEFT
CENTER    = WD_ALIGN_PARAGRAPH.CENTER

STYLE_ALIGN = {
    "Heading 1":   LEFT,
    "Heading 2":   LEFT,
    "Heading 3":   LEFT,
    "Normal":      JUSTIFY,
    "Body Text":   JUSTIFY,
    "Body Text 2": JUSTIFY,
    "Body Text 3": JUSTIFY,
}

# ── Files to process ─────────────────────────────────────────────────────────
REPO = Path(__file__).resolve().parents[2]

TARGETS = [
    "Critical Thinking of the Article/01079_Dzaki Muhammad Yusfian_Artikel 5.docx",
    "Critical Thinking of the Article/01079_Dzaki Muhammad Yusfian_Artikel 6.docx",
    "Critical Thinking of the Article/01079_Dzaki Muhammad Yusfian_Artikel 7.docx",
    "Critical Thinking of the Article/01079_Dzaki Muhammad Yusfian_Artikel 8.docx",
    "RMK/01079_Dzaki Muhammad Yusfian_RMK Pert. 4.docx",
    "RMK/01079_Dzaki Muhammad Yusfian_RMK Pert. 5.docx",
]


def _set_para_align_xml(para, alignment: WD_ALIGN_PARAGRAPH):
    """Set alignment directly on the paragraph XML so it overrides style."""
    pPr = para._p.get_or_add_pPr()
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    val_map = {
        WD_ALIGN_PARAGRAPH.JUSTIFY: "both",
        WD_ALIGN_PARAGRAPH.LEFT:    "left",
        WD_ALIGN_PARAGRAPH.CENTER:  "center",
        WD_ALIGN_PARAGRAPH.RIGHT:   "right",
    }
    jc.set(qn("w:val"), val_map[alignment])


def fix_document(rel_path: str) -> None:
    path = REPO / rel_path
    if not path.exists():
        print(f"SKIP (not found): {rel_path}")
        return

    doc = Document(str(path))

    # 1. Page size & margins
    for section in doc.sections:
        section.page_width    = A4_W
        section.page_height   = A4_H
        section.top_margin    = M_TOP
        section.bottom_margin = M_BOT
        section.left_margin   = M_LEFT
        section.right_margin  = M_RGT

    # 2. Paragraph alignment
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        target_align = STYLE_ALIGN.get(style_name)
        if target_align is not None:
            _set_para_align_xml(para, target_align)

    # 3. Also fix table cell paragraphs if any
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    style_name = para.style.name if para.style else ""
                    target_align = STYLE_ALIGN.get(style_name)
                    if target_align is not None:
                        _set_para_align_xml(para, target_align)

    doc.save(str(path))
    print(f"OK   {rel_path}")


def main():
    print(f"Formatting target: A4  |  Margins L{M_LEFT/_CM:.0f} R{M_RGT/_CM:.0f} T{M_TOP/_CM:.0f} B{M_BOT/_CM:.0f} cm")
    print(f"Alignment: Heading=Left  |  Normal/Body Text=Justify")
    print()
    for rel in TARGETS:
        fix_document(rel)
    print("\nDone.")


if __name__ == "__main__":
    main()
